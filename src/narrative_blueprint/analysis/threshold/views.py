"""Optional capped literal Excel and aligned Projector views."""

import os
import re

import numpy as np

from .contracts import AnalysisRefused
from .graph import read_tile
from .snapshot_input import checked_hash


class ViewFile:
    """Seekable output with byte admission before each write (including ZIP)."""

    def __init__(self, stream, exports, used):
        self.stream, self.exports, self.used = stream, exports, used
        self.high = 0

    def write(self, data):
        end = self.stream.tell() + len(data)
        high = max(end, self.high)
        limits = self.exports.scratch.budget.limits
        limits.check("max_view_bytes", self.used + high)
        limits.check("max_export_bytes", self.exports.bytes + self.used + high)
        self.exports.scratch.disk(high - self.high + 65536)
        self.high = high
        return self.stream.write(data)

    def __getattr__(self, name):
        return getattr(self.stream, name)


def literal_line(text):
    text = re.sub(r"[\t\r\n\x00-\x1f\x7f]", " ", text)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def excel_view(scratch, exports, count):
    from openpyxl import Workbook
    from openpyxl.cell import Cell

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Occurrences"
    sheet.append(["occurrence_id", "source_uuid", "text", "text_hash", "community_id", "outcome"])
    try:
        for row in scratch.db.execute("SELECT occurrence_id,source_uuid,text,text_hash,community,outcome FROM nodes ORDER BY node LIMIT ?", (count,)):
            cells = []
            for index, value in enumerate(row, 1):
                if isinstance(value, str) and len(value) > 32767:
                    raise AnalysisRefused("excel_cell_length", "Excel text would be truncated")
                cell = Cell(sheet, row=sheet.max_row + 1, column=index, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                cells.append(cell)
            sheet.append(cells)
            scratch.budget.tick()
        path = scratch.directory / "view.xlsx"
        with path.open("xb") as stream:
            workbook.save(ViewFile(stream, exports, 0))
            stream.flush()
            os.fsync(stream.fileno())
        return [path]
    finally:
        workbook.close()


def projector_view(scratch, exports, count, dimension):
    metadata = scratch.directory / "projector-metadata.tsv"
    vectors = scratch.directory / "projector-vectors.tsv"
    with metadata.open("xb") as meta_raw, vectors.open("xb") as vectors_raw, scratch.vectors.open("rb") as source:
        meta = ViewFile(meta_raw, exports, 0)
        data = ViewFile(vectors_raw, exports, 0)
        meta.write(b"occurrence_id\tsource_uuid\ttext\n")
        for node, occurrence, uuid, text in scratch.db.execute("SELECT node,occurrence_id,source_uuid,text FROM nodes ORDER BY node LIMIT ?", (count,)):
            meta.used = data.high
            meta.write(("\t".join(map(literal_line, (occurrence, uuid, text))) + "\n").encode("utf-8"))
            vector = read_tile(source, node, node + 1, dimension)[0]
            data.used = meta.high
            data.write(("\t".join(format(float(value), ".17g") for value in vector) + "\n").encode("utf-8"))
            scratch.budget.tick()
        for stream in (meta_raw, vectors_raw):
            stream.flush()
            os.fsync(stream.fileno())
    return [metadata, vectors]


def validate_view(name, paths, scratch, count, dimension):
    """Verify actual row counts and source alignment before listing a view."""
    if name == "excel":
        from openpyxl import load_workbook
        workbook = load_workbook(paths[0], read_only=True, data_only=False)
        try:
            actual = iter(workbook.active.iter_rows())
            next(actual)
            for expected in scratch.db.execute("SELECT occurrence_id,source_uuid,text,text_hash,community,outcome FROM nodes ORDER BY node LIMIT ?", (count,)):
                row = next(actual, ())
                if tuple(cell.value for cell in row) != expected or any(
                    cell.data_type == "f" for cell in row
                ):
                    raise AnalysisRefused("view_validation")
                scratch.budget.tick()
            if next(actual, None) is not None:
                raise AnalysisRefused("view_validation")
        finally:
            workbook.close()
        return
    with paths[0].open(encoding="utf-8") as metadata, paths[1].open(encoding="utf-8") as data, scratch.vectors.open("rb") as source:
        if metadata.readline() != "occurrence_id\tsource_uuid\ttext\n":
            raise AnalysisRefused("view_validation")
        for node, occurrence, uuid, text in scratch.db.execute("SELECT node,occurrence_id,source_uuid,text FROM nodes ORDER BY node LIMIT ?", (count,)):
            if metadata.readline().rstrip("\n").split("\t") != list(map(literal_line, (occurrence, uuid, text))):
                raise AnalysisRefused("view_validation")
            values = np.array([float(value) for value in data.readline().split("\t")])
            if not np.array_equal(values, read_tile(source, node, node + 1, dimension)[0]):
                raise AnalysisRefused("view_validation")
            scratch.budget.tick()
        if metadata.readline() or data.readline():
            raise AnalysisRefused("view_validation")


def write_views(scratch, exports, result, dimension, *, excel=False, projector=False):
    reports = {}
    count = min(result["nodes"], scratch.budget.limits.max_view_rows)
    for name, enabled, filenames in (("excel", excel, ["view.xlsx"]),
                                     ("projector", projector, ["projector-metadata.tsv", "projector-vectors.tsv"])):
        if not enabled:
            reports[name] = {"status": "not_requested"}
            continue
        report = {"rows": count, "total_nodes": result["nodes"], "truncated": count < result["nodes"],
                  "selection": "first canonical occurrence IDs", "files": []}
        try:
            text_bytes = scratch.db.execute('''SELECT coalesce(sum(length(CAST(text AS BLOB)) +
                length(CAST(occurrence_id AS BLOB)) + length(CAST(source_uuid AS BLOB))),0)
                FROM (SELECT * FROM nodes ORDER BY node LIMIT ?)''', (count,)).fetchone()[0]
            demand = text_bytes * 32 + count * 16384 + 16 * 1024**2
            from .contracts import process_memory
            scratch.budget.limits.check("max_working_bytes", process_memory()["current_bytes"] + demand)
            paths = excel_view(scratch, exports, count) if name == "excel" else projector_view(scratch, exports, count, dimension)
            validate_view(name, paths, scratch, count, dimension)
            for path in paths:
                size = path.stat().st_size
                exports.bytes += size
                exports.files[path.name] = {"bytes": size, "sha256": checked_hash(path, scratch.budget), "rows": count}
            report.update(status="complete", files=[p.name for p in paths])
        except Exception as exc:
            for filename in filenames:
                path = scratch.directory / filename
                if path.exists():
                    path.unlink()
                metadata = exports.files.pop(filename, None)
                if metadata:
                    exports.bytes -= metadata["bytes"]
            # A global time/memory/export/disk refusal must still stop publication.
            if isinstance(exc, AnalysisRefused) and exc.reason not in {"max_view_bytes", "excel_cell_length", "view_validation"}:
                raise
            report.update(status="refused" if isinstance(exc, AnalysisRefused) else "failed",
                          reason=getattr(exc, "reason", type(exc).__name__))
        reports[name] = report
    return reports
